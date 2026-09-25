from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUT = "ProtoGT Website Transition Plan.docx"
GREEN, DARK, PALE, BORDER = "087E3D", "16292B", "F3F5F4", "D9D9D9"

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    node = OxmlElement("w:shd")
    node.set(qn("w:fill"), fill)
    tcPr.append(node)

def border(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcPr.append(borders)
    for name in ("top", "left", "bottom", "right", "insideH", "insideV"):
        edge = OxmlElement("w:" + name)
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), "6")
        edge.set(qn("w:color"), BORDER)
        borders.append(edge)

def margins(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    m = OxmlElement("w:tcMar")
    for name, value in (("top", 100), ("start", 120), ("bottom", 100), ("end", 120)):
        n = OxmlElement("w:" + name)
        n.set(qn("w:w"), str(value))
        n.set(qn("w:type"), "dxa")
        m.append(n)
    tcPr.append(m)

def run(p, text, bold=False, color=None, size=None):
    r = p.add_run(text)
    r.bold = bold
    if color: r.font.color.rgb = RGBColor.from_string(color)
    if size: r.font.size = Pt(size)
    return r

def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(text)

def number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.add_run(text)

def table(doc, headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.autofit = False
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, text in enumerate(headers):
        c = t.rows[0].cells[i]
        c.width = Inches(widths[i]); shade(c, DARK); border(c); margins(c)
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        run(c.paragraphs[0], text, True, "FFFFFF", 9)
    for ri, values in enumerate(rows):
        cells = t.add_row().cells
        for i, text in enumerate(values):
            c = cells[i]
            c.width = Inches(widths[i]); border(c); margins(c)
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ri % 2: shade(c, PALE)
            c.paragraphs[0].paragraph_format.space_after = Pt(0)
            run(c.paragraphs[0], text, size=9)
    doc.add_paragraph()

doc = Document()
sec = doc.sections[0]
sec.top_margin, sec.bottom_margin = Inches(.72), Inches(.68)
sec.left_margin, sec.right_margin = Inches(.78), Inches(.78)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(10.5)
styles["Normal"].font.color.rgb = RGBColor.from_string(DARK)
styles["Normal"].paragraph_format.space_after = Pt(7)
styles["Normal"].paragraph_format.line_spacing = 1.12
for name, size, before, after in [("Title", 25, 0, 12), ("Heading 1", 16, 16, 7), ("Heading 2", 12, 11, 5)]:
    s = styles[name]
    s.font.name = "Aptos Display" if name == "Title" else "Aptos"
    s.font.size = Pt(size); s.font.bold = True; s.font.color.rgb = RGBColor(0, 0, 0)
    s.paragraph_format.space_before = Pt(before); s.paragraph_format.space_after = Pt(after)
if "Small Note" not in styles:
    s = styles.add_style("Small Note", WD_STYLE_TYPE.PARAGRAPH)
    s.font.name, s.font.size = "Aptos", Pt(9)
    s.font.color.rgb = RGBColor.from_string("555555")

footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
run(footer, "ProtoGT Website Transition Plan  |  Internal planning draft", color="666666", size=8)

doc.add_paragraph("ProtoGT Website Transition Plan", style="Title")
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(15)
run(p, "Moving the public website from Wix to a maintainable coded site", color=GREEN, size=13)
p = doc.add_paragraph()
run(p, "Purpose. ", True)
run(p, "This document sets out a practical route to replace the current Wix site with a fast, company-owned website. It covers the recommended stack, the transition process, changes to the site, and the customer, privacy, and operational features that should be included from launch.")
p = doc.add_paragraph()
run(p, "Recommended direction. ", True)
run(p, "Build a static Astro site with Content Collections for case studies, plain modern CSS for the design system, GitHub for source control, and Cloudflare Workers Static Assets for hosting. Keep Google Forms and Calendly for the first release, rather than building custom back-office systems too early.")

doc.add_paragraph("Scope", style="Heading 1")
for x in [
    "Preserve the current public information architecture, especially existing service and case-study URLs.",
    "Create a stronger, reusable case-study format that shows problem, constraints, delivery, and outcome.",
    "Retain existing quote and booking journeys while improving clarity, confirmation, and ownership of enquiries.",
    "Move hosting, DNS, source code, and operational access into company-controlled accounts.",
]: bullet(doc, x)

doc.add_paragraph("Working assumptions", style="Heading 2")
for x in [
    "The site is primarily a lead-generation and credibility site for engineering services and 3D printing, not a logged-in customer application.",
    "Case studies are published occasionally and can initially be maintained through a GitHub review workflow.",
    "The website should remain static at launch unless a genuine requirement for custom server functionality emerges.",
]: bullet(doc, x)

doc.add_page_break()
doc.add_paragraph("Recommended Technology", style="Heading 1")
p = doc.add_paragraph()
run(p, "Astro is the recommended framework. ", True)
run(p, "It matches a content-led, search-visible site and keeps browser JavaScript low by default. It also provides a clear structure for reusable pages and case-study content, while leaving room for small interactive components later.")
table(doc, ["Option", "Fit for ProtoGT", "Recommendation"], [
    ["Astro", "Content-first, fast static pages, structured case studies, optional interactive components.", "Recommended."],
    ["Eleventy", "Excellent lightweight static-site option with a simpler, more traditional template workflow.", "Viable alternative, but no clear advantage given existing Astro familiarity."],
    ["SvelteKit", "Strong if the site becomes more application-like or needs richer interactivity.", "More than needed for the initial release."],
    ["Next.js", "Mature React framework for full-stack applications.", "Use only if the team is already committed to Next.js or expects a substantial web app."],
    ["React with Vite", "Good tooling, but a client-rendered single-page app is not the best default for this SEO-led site.", "Not recommended for the main site."],
], [1.1, 3.7, 1.35])

doc.add_paragraph("Styling approach", style="Heading 2")
p = doc.add_paragraph()
run(p, "Start with plain modern CSS. ", True)
run(p, "Use CSS custom properties for ProtoGT colours, typography, spacing, responsive breakpoints, buttons, cards, and form treatments. Keep component styles alongside Astro components and reserve global styles for design tokens and base elements.")
bullet(doc, "Use Tailwind only if the people maintaining the site already prefer it and want a utility-first workflow.")
bullet(doc, "Avoid Bootstrap and large prebuilt component libraries; the site needs a distinct engineering-focused identity rather than a generic interface.")

doc.add_paragraph("Content model", style="Heading 2")
p = doc.add_paragraph()
run(p, "Use Astro Content Collections for case studies. ", True)
run(p, "Each case study should contain structured fields such as title, public client label, sector, services, hero image, short summary, project constraints, delivery, result, and SEO description. This is a content system within the repository, not a visual CMS. A visual editor can be added later if non-technical staff need to publish routinely.")

doc.add_page_break()
doc.add_paragraph("Site Changes", style="Heading 1")
table(doc, ["Area", "Change for the coded site"], [
    ["Homepage", "Clarify the primary audience, core services, proof points, and strongest next action: quote request or booked call."],
    ["Services", "Keep the three service areas clear and link each to relevant case studies and an appropriate enquiry route."],
    ["Case studies", "Use a consistent template: client problem, constraints, work delivered, evidence of outcome, relevant services, and next step."],
    ["Confidential work", "Use public-safe titles and summaries such as ‘Confidential industrial client’. Do not place private names in page source and cover them with CSS."],
    ["3D printing", "Make materials, process, turnaround, inspection, confidentiality, and the quote path immediately clear."],
    ["Contact journeys", "Keep Google Forms and Calendly initially, with clear confirmation, internal ownership, expected response time, and spam protection."],
    ["Performance", "Serve compressed, properly sized images; defer non-essential third-party embeds; test mobile performance before launch."],
], [1.45, 4.7])

doc.add_paragraph("Transition Plan", style="Heading 1")
p = doc.add_paragraph()
run(p, "The migration should be a parallel launch, not a rebuild performed on the live domain. ", True)
run(p, "Wix stays live while the Astro site is built and approved. DNS management and the public website cutover are treated as separate changes, which keeps the rollback path clear.")
table(doc, ["Phase", "What happens", "Exit criteria"], [
    ["1  Inventory", "Record every live page and URL, copy approved text, download original images, note forms, embeds, tracking, and page metadata.", "A page, asset, URL, and integration checklist exists."],
    ["2  Build", "Create the Astro site in a company GitHub repository. Build the homepage, service template, and one case study first, then complete remaining pages.", "A complete version runs on a non-production Cloudflare preview address."],
    ["3  Review", "Compare desktop and mobile views with Wix. Check copy, images, forms, booking, accessibility, links, redirects, and confidential-project treatment.", "Company sign-off and a launch checklist are complete."],
    ["4  DNS preparation", "Create a company Cloudflare account, copy existing DNS records, and verify mail records before changing nameservers. Keep the live web record serving Wix at this stage.", "The live Wix site and company email continue to work through Cloudflare DNS."],
    ["5  Cutover", "Attach protogt.co.uk and www.protogt.co.uk to the new Worker. Confirm SSL and routing, then switch the public web route during a quiet period.", "The production domain serves the new site and core journeys work."],
    ["6  Stabilise", "Monitor forms, booking, email, analytics, Search Console, broken links, and performance. Keep Wix accessible as a reference and rollback option.", "Two to four weeks of stable operation have passed."],
], [.85, 3.55, 1.75])
doc.add_paragraph("DNS and rollback safeguards", style="Heading 2")
for x in [
    "Before any nameserver change, copy and verify all DNS records, particularly MX, SPF, DKIM, DMARC, domain verification, and subdomain records.",
    "Do not cancel Wix at launch. Keep access for at least two to four weeks while the new site is monitored.",
    "Maintain a written record of the previous web-routing configuration and use Cloudflare deployment rollback if a production issue appears.",
    "Use company-owned Cloudflare and GitHub accounts, MFA, and at least two trusted administrators.",
]: bullet(doc, x)

doc.add_paragraph("Business Essentials", style="Heading 1")
p = doc.add_paragraph()
run(p, "This section is a planning checklist, not legal advice. ", True)
run(p, "The company director, accountant, and contract owner should confirm the final wording and obligations before publication.")
table(doc, ["Feature", "What the site should include"], [
    ["Company information", "If ProtoGT is a UK limited company, show the full registered name, company number, registered office, registration jurisdiction, and Limited or Ltd status in an easy-to-find footer or legal page."],
    ["Privacy notice", "Explain what personal data is collected through forms, booking, analytics, and enquiries; why it is used; the legal basis; retention; recipients or processors; rights; and contact route for data requests."],
    ["Cookie controls", "Audit the final site. Necessary cookies can run normally. Non-essential analytics, advertising, or third-party tracking should not load until the visitor gives a clear choice."],
    ["Forms and data", "Collect only what is needed, add spam protection, disclose the privacy notice beside the form, and define who receives and follows up each submission."],
    ["Terms", "Make approved quotation, payment, delivery, IP, confidentiality, warranty, and liability terms available before a customer accepts work."],
    ["Accessibility", "Target WCAG 2.2 AA in practice: keyboard use, visible focus, contrast, semantic headings, useful alt text, accessible forms, and descriptive links."],
    ["Trust information", "Publish only substantiated certifications, insurance, quality/testing practices, confidential-project handling, location, and practical lead-time information."],
], [1.45, 4.7])

doc.add_page_break()
doc.add_paragraph("Operating Model and Next Decisions", style="Heading 1")
doc.add_paragraph("Ownership", style="Heading 2")
for x in [
    "Name an owner for website content, technical maintenance, enquiry routing, privacy and cookie updates, and annual access review.",
    "Keep source code, assets, domain credentials, Cloudflare configuration, form ownership, Calendly ownership, analytics, and Search Console under company control.",
    "Use pull requests or an equivalent review step for public content changes, particularly case studies and legal pages.",
]: bullet(doc, x)
doc.add_paragraph("Measurement after launch", style="Heading 2")
for x in [
    "Track quote submissions, booked-call clicks, phone and email clicks, and case-study views.",
    "Review Search Console, form delivery, broken links, page speed, and mobile usability weekly for the first month.",
    "Use results to improve content and conversion paths in small iterations rather than overbuilding before launch.",
]: bullet(doc, x)
doc.add_paragraph("Decisions to confirm before build", style="Heading 2")
for x in [
    "Who is the primary website audience: engineering managers, product teams, manufacturing firms, 3D-print customers, or a mix?",
    "What is the primary conversion action for each key page: booked call, quote request, phone call, or email?",
    "Who will publish and approve future case studies, and is a GitHub-based workflow acceptable at launch?",
    "Which account will own the domain, Cloudflare, GitHub, Forms, Calendly, Search Console, and analytics?",
    "What privacy, cookie, company-registration, insurance, quality, and contractual information has the business approved for publication?",
]: number(doc, x)
doc.add_paragraph("First delivery milestone", style="Heading 2")
p = doc.add_paragraph()
run(p, "Create a short content and SEO inventory, then build the homepage and one complete case study as the design and content foundation. ", True)
run(p, "Once that direction is approved, the remaining service pages and case studies can be produced consistently and reviewed on Cloudflare preview deployments before any domain change.")

doc.save(OUT)
print(OUT)
